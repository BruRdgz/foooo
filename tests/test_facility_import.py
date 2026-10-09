import builtins
from dataclasses import asdict

import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from simgeope.world.config import DEFAULT_METRIC_CRS
from simgeope.world.exceptions import (
    NoRoadNearbyError,
    WorldBuildError,
    WorldDataMissingError,
    WorldFormatError,
)
from simgeope.world.facilities import Facility, FacilityIndex
from simgeope.world.manifest import sha256_file
from simgeope.world.sources.facilities import load_facilities


def import_source(path, world, **options):
    return load_facilities(
        path,
        world.roads,
        world.administrative,
        max_snap_distance_m=options.pop("max_snap_distance_m", 5000),
        **options,
    )


def write_csv(tmp_path, content):
    path = tmp_path / "pois.csv"
    path.write_text(content, encoding="utf-8-sig")
    return path


@pytest.mark.parametrize(
    "text",
    [
        "id,kind,lon,lat,name,municipality_code\nb,office,-34.98,-8,Serviço,9999999\na,hub,-35,-8,,2611606\n",
        " ID ; KIND ; Longitude ; Latitude ; Nome ; cod_municipio\nb;office;-34,98;-8;Serviço;9999999\na;hub;-35;-8;;2611606\n",
    ],
)
def test_csv_normalization_assignment_provenance_and_queries(
    tmp_path, tiny_world, text
):
    path = write_csv(tmp_path, text)
    index = import_source(path, tiny_world)
    assert [f.id for f in index.facilities] == ["a", "b"]
    assert index.get("b").name == "Serviço"
    assert index.get("b").municipality_code == "2611606"
    assert index.get("a").road_node == 10
    assert index.get("b").road_node == 30
    assert index.nearby((-35, -8), 100) == [index.get("a")]
    assert index.sources[0].filename == path.name
    assert index.sources[0].sha256 == sha256_file(path)
    assert len(index.warnings) == 1
    assert "coordinate-derived 2611606" in index.warnings[0]
    assert index.dropped_row_count == index.dropped_invalid_geometry_count == 0


def test_xlsx_legacy_farms_and_hubs_are_static_only(tmp_path, tiny_world):
    from openpyxl import Workbook

    path = tmp_path / "legacy.xlsx"
    book = Workbook()
    farm = book.active
    farm.title = "criadores"
    farm.append(
        [
            "id_criador",
            "nome",
            "longitude",
            "latitude",
            "total_aves",
            "biosseg_acesso",
            "infection",
            "status",
        ]
    )
    farm.append(["farm-a", "Criador", -35, -8, 45000, 0.6, "infected", "closed"])
    hubs = book.create_sheet("hubs")
    hubs.append(["id_hub", "tipo_hub", "nome", "longitude", "latitude", "n_veiculos"])
    hubs.append(["hub-a", "escritorio_tecnico", "Escritório", -34.98, -8, 3])
    instructions = book.create_sheet("instruções")
    instructions.append(["coluna", "descricao"])
    instructions.append(["total_aves", "dynamic scenario value"])
    book.save(path)
    book.close()

    index = import_source(path, tiny_world)
    assert index.get("farm-a").kind == "producer"
    assert index.get("hub-a").kind == "escritorio_tecnico"
    assert index.get("farm-a").name == "Criador"
    assert len(index.warnings) == 2
    assert "infection" in index.warnings[0]
    index.save(tmp_path / "built")
    table = pq.read_table(tmp_path / "built/facilities.parquet")
    assert {
        "total_aves",
        "biosseg_acesso",
        "infection",
        "status",
        "n_veiculos",
    }.isdisjoint(table.column_names)
    assert "infected" not in str(table.to_pylist())
    assert (
        FacilityIndex.load(tmp_path / "built", metric_crs=DEFAULT_METRIC_CRS).facilities
        == index.facilities
    )


def test_xlsx_canonical_schema_and_decimal_commas(tmp_path, tiny_world):
    from openpyxl import Workbook

    path = tmp_path / "canonical.xlsx"
    book = Workbook()
    book.active.title = "facilities"
    book.active.append(["id", "kind", "lon", "lat"])
    book.active.append(["a", "custom_kind", "-35,0", "-8,0"])
    book.save(path)
    book.close()
    assert import_source(path, tiny_world).get("a").road_node == 10


def test_parquet_recomputes_stale_snap_and_deterministic_persistence(
    tmp_path, tiny_world
):
    path = tmp_path / "input.parquet"
    pq.write_table(
        pa.Table.from_pylist(
            [
                {
                    "id": "b",
                    "kind": "hub",
                    "lon": -34.98,
                    "lat": -8.0,
                    "road_node": 999,
                    "name": "B",
                },
                {
                    "id": "a",
                    "kind": "farm",
                    "lon": -35.0,
                    "lat": -8.0,
                    "road_node": 999,
                    "name": None,
                },
            ]
        ),
        path,
    )
    index = import_source(path, tiny_world)
    assert index.get("a").road_node == 10
    assert index.get("b").road_node == 30
    index.save(tmp_path / "first")
    index.save(tmp_path / "second")
    assert sha256_file(tmp_path / "first/facilities.parquet") == sha256_file(
        tmp_path / "second/facilities.parquet"
    )
    assert (
        FacilityIndex.load(tmp_path / "first", metric_crs=DEFAULT_METRIC_CRS).facilities
        == index.facilities
    )


def test_old_canonical_facility_file_still_loads(tmp_path):
    facility = Facility("a", "farm", -35, -8)
    rows = [asdict(facility)]
    del rows[0]["name"]
    pq.write_table(pa.Table.from_pylist(rows), tmp_path / "facilities.parquet")
    assert (
        FacilityIndex.load(tmp_path, metric_crs=DEFAULT_METRIC_CRS).get("a") == facility
    )


@pytest.mark.parametrize(
    "change",
    [
        {"id": ""},
        {"kind": " "},
        {"lon": None},
        {"lon": float("nan")},
        {"lat": float("inf")},
        {"road_node": "10"},
        {"infection": "infected"},
        {"lon": "-35"},
        {"lat": True},
    ],
)
def test_invalid_canonical_persistence_is_refused(tmp_path, change):
    row = asdict(Facility("a", "farm", -35, -8)) | change
    pq.write_table(pa.Table.from_pylist([row]), tmp_path / "facilities.parquet")
    with pytest.raises(WorldFormatError, match="canonical facility"):
        FacilityIndex.load(tmp_path, metric_crs=DEFAULT_METRIC_CRS)


@pytest.mark.parametrize(
    "text,match",
    [
        (
            "id,kind,lon,lat\na,hub,-35,-8\na,hub,-35,-8\n",
            "row 3: duplicate facility ID",
        ),
        ("id,kind,lon,lat\n,hub,-35,-8\n", "row 2: empty facility ID"),
        ("id,kind,lon,lat\na,,-35,-8\n", "row 2: empty facility kind"),
        ("id,kind,lon,lat\na,hub,,-8\n", "row 2: lon"),
        ("id,kind,lon,lat\na,hub,-35,nan\n", "row 2: lat must be finite"),
        ("id,kind,lon,lat\na,hub,inf,-8\n", "row 2: lon must be finite"),
        ("id,kind,lon,lat\na,hub,-181,-8\n", "row 2: lon must be finite"),
        ("id,kind,lon,lat\na,hub,-35,91\n", "row 2: lat must be finite"),
        ("id,kind,lon,lat\na,hub,text,-8\n", "row 2: invalid lon"),
        ("id,kind,lon,lat\na,hub,-35,-8\n\n", "row 3: blank facility row"),
        ("id,kind,lon,lat\na,hub,-35,-8,extra\n", "row 2: row has values beyond"),
        ("id,kind,latitude\na,hub,-8\n", "row 1: missing required columns: lon"),
        (
            "id,kind,lon,longitude,lat\na,hub,-35,-35,-8\n",
            "row 1: duplicate or ambiguous",
        ),
        ("id,kind,lon,lat\n", "row 2: facility table has no data rows"),
    ],
)
def test_csv_row_errors_have_file_and_line(tmp_path, tiny_world, text, match):
    path = write_csv(tmp_path, text)
    with pytest.raises(WorldBuildError, match=match) as error:
        import_source(path, tiny_world)
    assert str(path) in str(error.value)


def test_xlsx_duplicate_ids_and_unknown_sheet_have_sheet_diagnostics(
    tmp_path, tiny_world
):
    from openpyxl import Workbook

    path = tmp_path / "duplicate.xlsx"
    book = Workbook()
    for sheet in [book.active, book.create_sheet("second")]:
        sheet.append(["id", "kind", "lon", "lat"])
        sheet.append(["a", "farm", -35, -8])
    book.save(path)
    with pytest.raises(WorldBuildError, match="sheet 'second'.*row 2: duplicate"):
        import_source(path, tiny_world)
    book["second"].delete_rows(1, 2)
    book["second"].append(["unknown", "columns"])
    book.save(path)
    book.close()
    with pytest.raises(WorldBuildError, match="sheet 'second'.*row 1: missing"):
        import_source(path, tiny_world)


def test_outside_municipality_clears_stale_code(tmp_path, tiny_world):
    path = write_csv(
        tmp_path, "id,kind,lon,lat,municipality_code\na,office,-35.015,-8,2611606\n"
    )
    index = import_source(path, tiny_world)
    assert index.get("a").municipality_code is None
    assert index.get("a").road_node == 10
    assert "outside administrative municipalities" in index.warnings[0]


def test_snap_limit_is_contextual_and_snap_occurs_once(
    tmp_path, tiny_world, monkeypatch
):
    import simgeope.world.sources.facilities as module

    path = write_csv(tmp_path, "id,kind,lon,lat\na,office,-35,-8.001\n")
    with pytest.raises(
        NoRoadNearbyError, match="row 2: Facility 'a'.*snapped road node"
    ):
        import_source(path, tiny_world, max_snap_distance_m=1)
    calls = []
    original = module.snap_facility

    def spy(*args, **kwargs):
        calls.append(args[0].id)
        return original(*args, **kwargs)

    monkeypatch.setattr(module, "snap_facility", spy)
    with pytest.warns(UserWarning, match="row 2: Facility 'a'"):
        index = import_source(
            path, tiny_world, max_snap_distance_m=1, allow_excessive_snap=True
        )
    assert calls == ["a"]
    assert 100 < index.get("a").snap_distance_m < 120
    assert "snapped road node" in index.warnings[0]


def test_missing_spreadsheet_dependency_is_actionable(
    tmp_path, tiny_world, monkeypatch
):
    path = tmp_path / "input.xlsx"
    path.touch()
    original_import = builtins.__import__

    def missing(name, *args, **kwargs):
        if name == "openpyxl":
            raise ModuleNotFoundError("No module named 'openpyxl'")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", missing)
    with pytest.raises(
        WorldBuildError, match=r"optional spreadsheet extra.*simgeope\[spreadsheet\]"
    ):
        import_source(path, tiny_world)


def test_coordinate_outside_metric_projection_has_row_diagnostic(tmp_path, tiny_world):
    path = write_csv(tmp_path, "id,kind,lon,lat\na,office,52.5,0\n")
    with pytest.raises(WorldBuildError, match="row 2: cannot snap facility 'a'"):
        import_source(path, tiny_world, allow_excessive_snap=True)


def test_source_missing_unsupported_and_corrupt(tmp_path, tiny_world):
    with pytest.raises(WorldDataMissingError, match="Facility source missing"):
        import_source(tmp_path / "missing.csv", tiny_world)
    path = tmp_path / "input.xls"
    path.touch()
    with pytest.raises(WorldBuildError, match="unsupported facility format"):
        import_source(path, tiny_world)
    path = tmp_path / "corrupt.xlsx"
    path.write_text("not a workbook")
    with pytest.raises(WorldBuildError, match="Cannot read facility source"):
        import_source(path, tiny_world)
    path = write_csv(tmp_path, "")
    with pytest.raises(WorldBuildError, match="empty facility source"):
        import_source(path, tiny_world)
