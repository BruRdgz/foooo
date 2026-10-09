"""Local static POI import; legacy simulation columns are never retained.

The canonical input columns are id, kind, lon, lat (WGS84), and optional
name/municipality_code. CSV accepts comma or semicolon delimiters and decimal
commas. XLSX also accepts the legacy criadores/hubs sheets. All rows must be
valid; blank rows and duplicate identifiers are errors, not dropped records.
"""

from __future__ import annotations

import csv
import math
import unicodedata
import warnings
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import pyarrow as pa
import pyarrow.parquet as pq

from ..exceptions import NoRoadNearbyError, WorldBuildError, WorldDataMissingError
from ..facilities import Facility, FacilityIndex, snap_facility
from ..manifest import SourceRecord

_ALIASES = {
    "longitude": "lon",
    "latitude": "lat",
    "nome": "name",
    "cod_municipio": "municipality_code",
    "codigo_municipio": "municipality_code",
    "id_criador": "id",
    "id_hub": "id",
    "tipo_hub": "kind",
}
_STATIC_COLUMNS = {"id", "kind", "lon", "lat", "name", "municipality_code"}
_INSTRUCTION_SHEETS = {"instrucoes", "instructions"}


def _label(value):
    value = unicodedata.normalize("NFKD", str(value).strip().casefold())
    value = "".join(c for c in value if not unicodedata.combining(c))
    return value.replace(" ", "_").replace("-", "_")


def _empty(value):
    return value is None or isinstance(value, str) and not value.strip()


def _text(value):
    if _empty(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _location(path, sheet, row):
    return f"{path}" + (f" [sheet {sheet!r}]" if sheet else "") + f" row {row}"


def _schema(headers, context):
    headers = list(headers)
    while headers and _empty(headers[-1]):
        headers.pop()
    if not headers or any(_empty(header) for header in headers):
        raise WorldBuildError(f"{context}: missing/blank column headers")
    original = [_label(header) for header in headers]
    normalized = [_ALIASES.get(header, header) for header in original]
    if len(set(normalized)) != len(normalized):
        raise WorldBuildError(f"{context}: duplicate or ambiguous column headers")
    legacy_producer = "id_criador" in original
    required = {"id", "lon", "lat"} | (set() if legacy_producer else {"kind"})
    missing = required - set(normalized)
    if missing:
        raise WorldBuildError(
            f"{context}: missing required columns: {', '.join(sorted(missing))}"
        )
    return normalized, "producer" if legacy_producer else None


def _csv_tables(path):
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(8192)
        stream.seek(0)
        if not sample.strip():
            raise WorldBuildError(f"{path}: empty facility source")
        # Restrict detection to supported separators, avoiding decimal punctuation.
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(stream, dialect=dialect, strict=True)
        headers = next(reader, ())
        yield None, headers, ((reader.line_num, row) for row in reader)


def _xlsx_tables(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise WorldBuildError(
            f"{path}: XLSX import requires the optional spreadsheet extra; "
            "install 'simgeope[spreadsheet]' or use CSV/Parquet"
        ) from exc
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        found = False
        for sheet in book:
            if _label(sheet.title) in _INSTRUCTION_SHEETS:
                continue
            found = True
            rows = sheet.iter_rows(values_only=True)
            yield sheet.title, next(rows, ()), enumerate(rows, start=2)
        if not found:
            raise WorldBuildError(f"{path}: no facility data sheets")
    finally:
        book.close()


def _parquet_tables(path):
    table = pq.read_table(path)
    yield (
        None,
        table.column_names,
        (
            (row_number, [row[column] for column in table.column_names])
            for row_number, row in enumerate(table.to_pylist(), start=2)
        ),
    )


def _coordinate(value, field, context):
    if isinstance(value, bool) or _empty(value):
        raise WorldBuildError(f"{context}: {field} must be a finite WGS84 coordinate")
    try:
        number = float(
            value.strip().replace(",", ".") if isinstance(value, str) else value
        )
    except (TypeError, ValueError, OverflowError) as exc:
        raise WorldBuildError(
            f"{context}: invalid {field} coordinate {value!r}"
        ) from exc
    limit = 180 if field == "lon" else 90
    if not math.isfinite(number) or not -limit <= number <= limit:
        raise WorldBuildError(
            f"{context}: {field} must be finite and within "
            f"WGS84 range [-{limit}, {limit}]"
        )
    return number


def load_facilities(
    path,
    roads,
    administrative,
    *,
    max_snap_distance_m,
    allow_excessive_snap=False,
) -> FacilityIndex:
    """Validate, assign municipality and snap each record exactly once.

    The returned index exposes sources (SourceRecord tuple), warnings (tuple),
    dropped_row_count and dropped_invalid_geometry_count (both zero). Invalid
    inputs fail atomically; persistence is performed by FacilityIndex.save().
    Municipality is always derived from administrative geometry, even if an
    input code is stale. Outside points retain None and a diagnostic warning.
    """
    path = Path(path)
    if not path.is_file():
        raise WorldDataMissingError(f"Facility source missing: {path}")
    if not math.isfinite(max_snap_distance_m) or max_snap_distance_m < 0:
        raise ValueError(
            "Facility maximum snap distance must be finite and nonnegative"
        )
    loaders = {".csv": _csv_tables, ".xlsx": _xlsx_tables, ".parquet": _parquet_tables}
    loader = loaders.get(path.suffix.lower())
    if loader is None:
        raise WorldBuildError(
            f"{path}: unsupported facility format; use CSV, XLSX or Parquet"
        )
    if administrative is None:
        raise ValueError("Facility import requires an administrative layer")
    facilities, messages, seen = [], [], {}
    try:
        for sheet, headers, rows in loader(path):
            context = _location(path, sheet, 1)
            columns, default_kind = _schema(headers, context)
            ignored = sorted(set(columns) - _STATIC_COLUMNS)
            if ignored:
                shown = ", ".join(ignored[:12])
                remaining = f" (+{len(ignored) - 12} more)" if len(ignored) > 12 else ""
                messages.append(
                    f"{context}: discarded non-static/unsupported columns: "
                    f"{shown}{remaining}"
                )
            count = 0
            for row_number, values in rows:
                context = _location(path, sheet, row_number)
                values = list(values)
                if all(_empty(value) for value in values):
                    raise WorldBuildError(f"{context}: blank facility row")
                if any(not _empty(value) for value in values[len(columns) :]):
                    raise WorldBuildError(
                        f"{context}: row has values beyond its column headers"
                    )
                values = values[: len(columns)] + [None] * max(
                    0, len(columns) - len(values)
                )
                row = dict(zip(columns, values))
                facility_id = _text(row["id"])
                if not facility_id:
                    raise WorldBuildError(f"{context}: empty facility ID")
                if facility_id in seen:
                    raise WorldBuildError(
                        f"{context}: duplicate facility ID {facility_id!r}; "
                        f"first seen at {seen[facility_id]}"
                    )
                seen[facility_id] = context
                kind = _text(row.get("kind")) or default_kind
                if not kind:
                    raise WorldBuildError(f"{context}: empty facility kind")
                facility = Facility(
                    facility_id,
                    kind,
                    _coordinate(row["lon"], "lon", context),
                    _coordinate(row["lat"], "lat", context),
                    municipality_code=_text(row.get("municipality_code")),
                    name=_text(row.get("name")),
                )
                try:
                    with warnings.catch_warnings(record=True) as captured:
                        warnings.simplefilter("always", UserWarning)
                        snapped = snap_facility(
                            facility,
                            roads,
                            administrative,
                            max_distance_m=max_snap_distance_m,
                            allow_excessive=allow_excessive_snap,
                        )
                except NoRoadNearbyError as exc:
                    raise NoRoadNearbyError(f"{context}: {exc}") from exc
                except (ValueError, TypeError) as exc:
                    raise WorldBuildError(
                        f"{context}: cannot snap facility {facility_id!r}: {exc}"
                    ) from exc
                for warning in captured:
                    message = f"{context}: {warning.message}"
                    messages.append(message)
                    warnings.warn(message, UserWarning, stacklevel=2)
                if snapped.municipality_code is None:
                    messages.append(
                        f"{context}: facility {facility_id!r} is outside "
                        "administrative municipalities; municipality_code cleared"
                    )
                elif (
                    facility.municipality_code
                    and facility.municipality_code != snapped.municipality_code
                ):
                    messages.append(
                        f"{context}: facility {facility_id!r} municipality_code "
                        f"replaced by coordinate-derived {snapped.municipality_code}"
                    )
                facilities.append(snapped)
                count += 1
            if not count:
                raise WorldBuildError(
                    f"{_location(path, sheet, 2)}: facility table has no data rows"
                )
    except (
        OSError,
        UnicodeError,
        csv.Error,
        BadZipFile,
        ParseError,
        pa.ArrowException,
    ) as exc:
        raise WorldBuildError(f"Cannot read facility source {path}: {exc}") from exc
    index = FacilityIndex(facilities, metric_crs=roads.metric_crs)
    index.sources = (SourceRecord.from_path(path),)
    index.warnings = tuple(messages)
    return index
